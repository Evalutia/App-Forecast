#!/usr/bin/env python3
"""Genera docs/Planilla_Reposicion_Guia_Cliente.docx a partir del contenido acordado
en la sesión /grill-me del 2026-06-20. Re-ejecutar este script regenera el .docx desde
cero (no editar el .docx a mano si se va a volver a correr este script)."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Estilos base ────────────────────────────────────────────────────────────
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def formula_box(text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph()

def nota(text):
    p = doc.add_paragraph()
    run = p.add_run("Nota: " + text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# ── Portada ──────────────────────────────────────────────────────────────────
title = doc.add_heading("Guía de interpretación — Planilla de Reposición", level=0)
sub = doc.add_paragraph()
sub_run = sub.add_run("Evalutia · Documento de revisión técnica · 20 de junio de 2026")
sub_run.italic = True
sub_run.font.size = Pt(11)
sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# ── 1. Objetivo ──────────────────────────────────────────────────────────────
doc.add_heading("1. Objetivo de este documento", level=1)
doc.add_paragraph(
    "Este documento explica cómo se calcula y cómo se debe leer la Planilla de "
    "Reposición: qué datos usa, qué significa cada columna y, en particular, qué "
    "significa el color de cada celda. No es un manual de uso de la pantalla — es la "
    "explicación de la lógica de negocio detrás de los números."
)

# ── 2. Qué es ────────────────────────────────────────────────────────────────
doc.add_heading("2. ¿Qué es la Planilla de Reposición?", level=1)
doc.add_paragraph(
    "Es una tabla con un renglón por producto (SKU), que cruza 13 meses de venta "
    "(el mes en curso más los 12 anteriores) con el stock que tuvo disponible cada "
    "uno de esos meses. El objetivo es responder dos preguntas para cada producto: "
    "¿qué tan bien rota? y ¿cuántos días quedan antes de quedarse sin stock?"
)

# ── 3. Datos que usa ─────────────────────────────────────────────────────────
doc.add_heading("3. Qué datos usa", level=1)
doc.add_paragraph("Para cada producto, el sistema combina cuatro fuentes de información:")
for item in [
    "Unidades vendidas por mes.",
    "Stock diario, sumado entre todos los depósitos.",
    "Stock mínimo configurado por artículo (el umbral debajo del cual se considera "
    "que el producto no tiene stock disponible para vender con normalidad, aunque "
    "el sistema registre alguna unidad).",
    "Factor estacional mensual por producto, para poder corregir el efecto de "
    "estacionalidad (ej.: ventas más altas en diciembre) al calcular la rotación.",
]:
    doc.add_paragraph(item, style="List Bullet")

# ── 4. Conceptos base ────────────────────────────────────────────────────────
doc.add_heading("4. Conceptos base", level=1)

doc.add_heading("4.1 «Día con stock»", level=2)
doc.add_paragraph(
    "Un día cuenta como «día con stock» cuando, sumando todos los depósitos, las "
    "unidades disponibles superan el stock mínimo configurado para ese artículo. Si "
    "un producto tiene 2 unidades pero su stock mínimo es 5, ese día NO se considera "
    "un día con stock real para los cálculos de rotación."
)
formula_box("día con stock  ⟺  stock total del día  >  stock mínimo del artículo")

doc.add_heading("4.2 Mes cerrado vs. mes en curso", level=2)
doc.add_paragraph(
    "La planilla siempre mira 13 meses: el mes actual (en curso) y los 12 anteriores "
    "ya cerrados. Esta distinción importa porque el mes en curso todavía no terminó: "
    "si se le aplicaran los mismos criterios que a un mes cerrado, mostraría un "
    "porcentaje de días sin stock artificialmente alto solo por estar a mitad de "
    "mes. El tratamiento especial del mes en curso se explica en la sección 5."
)

# ── 5. Estado del mes ────────────────────────────────────────────────────────
doc.add_heading("5. Estado del mes: Normal / Quiebre parcial / Sin stock", level=1)
doc.add_paragraph(
    "Para cada combinación producto + mes cerrado, el sistema clasifica el mes en "
    "uno de tres estados, según la proporción de días con stock sobre el total de "
    "días del mes:"
)
for item in [
    "Normal: tuvo stock disponible el 100% de los días del mes (ningún día de "
    "quiebre).",
    "Quiebre parcial: tuvo stock al menos un día, pero no todos. Un solo día sin "
    "stock ya alcanza para salir de «Normal» — no hay tolerancia mínima.",
    "Sin stock: no tuvo stock disponible ningún día del mes.",
]:
    doc.add_paragraph(item, style="List Bullet")
formula_box(
    "días con stock = 0                              →  Sin stock\n"
    "días con stock / días del mes  ≥  100%           →  Normal\n"
    "cualquier otro caso                              →  Quiebre parcial"
)
nota(
    "el criterio de «100%, sin tolerancia» replica exactamente el criterio que el "
    "cliente ya usaba en su propia planilla manual — no es un umbral arbitrario "
    "nuevo del sistema, es el mismo que el cliente ya aplicaba a ojo."
)
doc.add_paragraph()
doc.add_paragraph(
    "Caso especial — mes en curso: como el mes en curso no terminó, no se le aplica "
    "el umbral del 100%. Solo se evalúa si tuvo o no stock hasta la fecha: si tuvo "
    "al menos un día con stock, se muestra «Normal»; si no tuvo ningún día con "
    "stock, se muestra «Sin stock». Esto evita que, por ejemplo, el día 5 del mes ya "
    "se pinte de quiebre solo porque «faltan» 25 días para llegar al 100%."
)

# ── 6. Frecuencia ────────────────────────────────────────────────────────────
doc.add_heading("6. Frecuencia de venta del producto: Alta / Media / Baja", level=1)
doc.add_paragraph(
    "No todos los productos venden con la misma regularidad. Antes de evaluar un "
    "quiebre, el sistema clasifica cada producto según en cuántos de los últimos "
    "12 meses cerrados tuvo al menos una venta:"
)
for item in [
    "Alta frecuencia: vendió en 9 o más de los últimos 12 meses.",
    "Media frecuencia: vendió en 4 a 8 de los últimos 12 meses.",
    "Baja frecuencia: vendió en 3 o menos de los últimos 12 meses.",
]:
    doc.add_paragraph(item, style="List Bullet")
nota(
    "estos umbrales (9 / 4–8 / 3 meses) son un parámetro configurable del sistema, "
    "no una regla matemática fija. Están pendientes de validación final con el "
    "equipo económico del cliente y pueden ajustarse."
)

# ── 7. Por qué cambia la fórmula ─────────────────────────────────────────────
doc.add_heading("7. Por qué la fórmula de rotación cambia según la frecuencia", level=1)
doc.add_paragraph(
    "La «rotación diaria» mide, en promedio, cuántas unidades se venden por día. "
    "Cuando un mes tuvo quiebre parcial, hay que decidir sobre qué base de días "
    "calcular ese promedio — y ahí la frecuencia del producto importa:"
)
doc.add_paragraph(
    "Alta frecuencia: rotación = ventas del mes ÷ días con stock. Como el producto "
    "vende seguido, los días con stock real son una muestra suficientemente grande "
    "y confiable para calcular la rotación real."
)
formula_box("rotación  =  ventas ÷ días con stock")
doc.add_paragraph(
    "Baja frecuencia: rotación = ventas del mes ÷ días naturales del mes "
    "(calendario completo). Un producto que vende poco tiene muy pocos días de "
    "venta real en el mes — si se dividiera solo por esos pocos días, un único "
    "pico de venta inflaría artificialmente la rotación. Dividir por el mes "
    "completo «diluye» ese ruido y da una estimación más conservadora y estable."
)
formula_box("rotación  =  ventas ÷ días naturales del mes")
doc.add_paragraph("Media frecuencia: promedio de ambas fórmulas.")
formula_box("rotación  =  ( ventas/días con stock  +  ventas/días naturales del mes ) ÷ 2")
doc.add_paragraph(
    "En resumen: cuanto menos vende un producto, menos confiamos en sus pocos días "
    "de venta real para proyectar — por eso se «estira» el cálculo sobre todo el "
    "mes en vez de concentrarlo solo en los días con stock."
)

# ── 8. Celdas coloreadas ─────────────────────────────────────────────────────
doc.add_heading("8. Las celdas coloreadas: cómo leerlas", level=1)
doc.add_paragraph(
    "Este es el punto central de la planilla. Cada celda mensual de ventas y de "
    "rotación se pinta según el estado del mes y, si hubo quiebre, según la "
    "frecuencia del producto:"
)

legend = [
    ("Amarillo", "FFCA28", "Quiebre parcial en un producto de ALTA frecuencia"),
    ("Naranja",  "FFB74D", "Quiebre parcial en un producto de MEDIA frecuencia"),
    ("Rojo",     "EF9A9A", "Quiebre parcial en un producto de BAJA frecuencia"),
    ("Gris",     "90A4AE", "Sin stock durante todo el mes"),
    ("Sin color","FFFFFF", "Mes normal — sin quiebre"),
]
table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Color"
hdr[1].text = "Significado"
for color_name, hex_color, meaning in legend:
    row = table.add_row().cells
    set_cell_shading(row[0], hex_color)
    row[0].text = color_name
    row[1].text = meaning
doc.add_paragraph()

doc.add_paragraph(
    "¿Por qué el mismo «quiebre parcial» se pinta distinto según la frecuencia? La "
    "intensidad del color escala con el nivel de incertidumbre del dato: en un "
    "producto de alta frecuencia, un quiebre parcial es una señal confiable — hay "
    "muchos días de venta real respaldando el número — y se mantiene en amarillo, "
    "el mismo color que el cliente ya usaba en su planilla original para señalar "
    "quiebre. A medida que baja la frecuencia, hay menos datos reales detrás del "
    "número y el quiebre es más difícil de diagnosticar con precisión — el color "
    "se intensifica (naranja, luego rojo) para llamar la atención sobre esas filas, "
    "no porque matemáticamente sea «peor», sino porque amerita revisión manual más "
    "cuidadosa."
)

# ── 9. Columnas de resumen ───────────────────────────────────────────────────
doc.add_heading("9. Columnas de resumen y sugerencia de reposición", level=1)

doc.add_heading("VTA — Ventas totales", level=2)
doc.add_paragraph(
    "Suma de unidades vendidas en los 12 meses cerrados (no incluye el mes en "
    "curso, que está incompleto)."
)
formula_box("VTA  =  suma de ventas de los 12 meses cerrados")

doc.add_heading("DDSTK — Demanda diaria con stock", level=2)
doc.add_paragraph(
    "Rotación diaria promedio de los 13 meses, calculada como el total de unidades "
    "vendidas dividido por el total de días con stock disponible en esos 13 meses. "
    "A diferencia de ROT.S, no pondera por antigüedad ni distingue frecuencia."
)
formula_box("DDSTK  =  (suma de ventas) ÷ (suma de días con stock), de los 13 meses")

doc.add_heading("Rot. DesEstac. — Rotación desestacionalizada", level=2)
doc.add_paragraph(
    "Promedio de la rotación diaria de los meses normales y con quiebre parcial, "
    "corrigiendo el efecto de estacionalidad (un producto que vende mucho en "
    "diciembre no debe parecer «más rotativo» todo el año solo por ese pico). En "
    "los meses con quiebre parcial usa la rotación ajustada por frecuencia "
    "(sección 7) en vez de la rotación cruda. Excluye el mes en curso."
)

doc.add_heading("ROT.S — Rotación sugerida", level=2)
doc.add_paragraph(
    "Es el número que el sistema recomienda usar para proyectar reposición. Es un "
    "promedio de los últimos 13 meses, pero no es un promedio simple: los meses "
    "más recientes pesan más que los meses más viejos (peso lineal — el mes más "
    "reciente pesa 13 veces más que el más antiguo de la ventana). En los meses "
    "con quiebre parcial usa la rotación ajustada por frecuencia, no la rotación "
    "cruda. Requiere al menos 3 meses con datos utilizables; si no los tiene, "
    "muestra «—»."
)
formula_box(
    "ROT.S  =  Σ (peso del mes × rotación del mes) ÷ Σ (peso del mes)\n"
    "peso = 13 para el mes más reciente, 1 para el más antiguo de la ventana"
)

doc.add_heading("Fiabilidad %", level=2)
doc.add_paragraph(
    "Mide qué tan estable fue la rotación mes a mes — no qué tan alta o baja es. "
    "Un producto que vende siempre una cantidad parecida cada mes tiene fiabilidad "
    "alta: ROT.S es un buen predictor. Un producto con ventas muy irregulares (un "
    "mes 50 unidades, el siguiente 2) tiene fiabilidad baja: el número de ROT.S "
    "existe, pero hay que tomarlo con cautela."
)
nota("colores — Verde ≥ 70% (confiable) · Amarillo 40–69% (parcial) · Rojo < 40% (poco confiable).")

doc.add_heading("QBK — Días hasta quiebre", level=2)
doc.add_paragraph(
    "Con el stock actual y la rotación sugerida (ROT.S), estima en cuántos días se "
    "agotaría el stock si la venta sigue al ritmo actual."
)
formula_box("QBK  =  stock actual ÷ ROT.S")
nota(
    "colores — Rojo = 0 días (ya está en quiebre) · Amarillo ≤ 15 días (margen "
    "ajustado — 15 días es el lead time típico de reposición) · Verde > 15 días "
    "(margen suficiente)."
)

# ── 10. Glosario ─────────────────────────────────────────────────────────────
doc.add_heading("10. Glosario de abreviaturas", level=1)
gloss_table = doc.add_table(rows=1, cols=2)
gloss_table.style = "Light Grid Accent 1"
hdr = gloss_table.rows[0].cells
hdr[0].text = "Sigla"
hdr[1].text = "Significado"
for sigla, significado in [
    ("SKU", "Código único de producto."),
    ("VTA", "Ventas totales de los 12 meses cerrados."),
    ("DDSTK", "Demanda diaria con stock."),
    ("Rot. DesEstac.", "Rotación diaria desestacionalizada."),
    ("ROT.S", "Rotación sugerida."),
    ("QBK", "Días hasta quiebre."),
]:
    row = gloss_table.add_row().cells
    row[0].text = sigla
    row[1].text = significado
doc.add_paragraph()

# ── 11. Parámetros ajustables ────────────────────────────────────────────────
doc.add_heading("11. Parámetros sujetos a ajuste", level=1)
doc.add_paragraph(
    "Estos valores no son constantes matemáticas — son criterios de negocio "
    "configurables, definidos para esta primera versión y revisables a futuro:"
)
for item in [
    "Umbral de «mes normal»: 100% de días con stock, sin tolerancia.",
    "Umbrales de frecuencia: 9 / 4–8 / 3 meses (alta / media / baja).",
    "Lead time de referencia para QBK: 15 días.",
    "Mínimo de meses con datos para calcular ROT.S: 3 meses.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.save(r"c:\Users\Nico\Evalutia\App-Forecast\docs\Planilla_Reposicion_Guia_Cliente.docx")
print("Documento generado.")
