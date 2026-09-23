#!/usr/bin/env python3
"""
cargador_archivos_cliente.py — Issue #195.

Cargador de los dos archivos que el cliente entregó en septiembre de 2026, y
comparación de cobertura de catálogo contra `articulos`. Es el prefactor del
diagnóstico de calidad de datos: los issues #196 a #201 consumen este módulo,
así que las funciones de acá son la definición operativa de cómo se leen esos
archivos -- si una regla de parseo cambia, cambia para todo el diagnóstico.

Solo lectura: no escribe nada en la base ni toca los archivos.

Requiere `openpyxl` además de `pymysql` (ver services/etl/requirements-dev.txt).
La imagen del contenedor etl NO trae openpyxl, así que la forma práctica de
correrlo es desde el host, con un túnel al MySQL de la VM:

  ssh -f -N -L 13307:localhost:3307 <vm>
  MYSQL_HOST=127.0.0.1 MYSQL_PORT=13307 MYSQL_USER=... MYSQL_PASSWORD=... \
    python3 scripts/cargador_archivos_cliente.py

Con los archivos en otra ubicación:
  ARCHIVOS_CLIENTE_DIR=/ruta python3 scripts/cargador_archivos_cliente.py

Sin credenciales corre igual y reporta solo lo que sale de los archivos.

Las dos trampas que motivaron este módulo, ambas verificadas sobre los archivos
reales:

1. **La regla nulo/cero es asimétrica.** En las columnas de venta el cliente usa
   celda vacía para "no vendió" y nunca 0 (0 ceros explícitos en 70.564 celdas);
   en las de rotación mensual, que están densas, el cero es un cero real.
   Aplicar una sola regla a las dos familias genera miles de falsos positivos.
   Es el problema de "cero implícito" de #187/#188, en espejo.

2. **El archivo de movimientos no es una tabla plana**, es un reporte bandeado
   con 35 filas de subtotal, una fila `Total General` con las columnas corridas
   y un pie de página con número de página numérico. Un parser que no las filtre
   triplica los totales: durante la verificación del plan dio +1209/-2457/+1399
   en vez de +403/-819/+466.
"""

import os
import sys
from dataclasses import dataclass

# openpyxl se importa dentro de las funciones de lectura, no acá: las funciones
# puras de este módulo (y sus tests) no necesitan leer Excel, y la imagen del
# contenedor etl no lo trae. Un import de módulo rompería la colección de toda
# la suite en un entorno sin la librería. Mismo criterio que #127.

DIR_POR_DEFECTO = os.environ.get(
    "ARCHIVOS_CLIENTE_DIR",
    os.path.join(os.path.dirname(os.path.abspath(__file__)),
                 "fixtures", "archivos_cliente_2026-09"))

ARCHIVO_VENTAS = "ventas_hasta_2026-08_generos-todos.xlsx"
ARCHIVO_MOVIMIENTOS = "mov_stock_deposito-5_2026-07.xlsx"

HOJA_VENTAS = "Ventas"
HOJA_MOVIMIENTOS = "Sheet1"

MESES_ES = {m: i for i, m in enumerate(
    ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"], start=1)}

COL_ROT_DESESTAC = "Rotacion DesEstac."
COL_ROT_MANUAL = "Rot. Manual"
COL_ESTADO = "Estado Art."
COL_DESCRIPCION = "Descripcion"
COL_BARCODE = "Codigos Barras"

ENTRADA_MERCADERIA = "Entrada de Mercaderia"
AJUSTE_SALIDA = "Ajuste de Salida"
AJUSTE_ENTRADA = "Ajuste de Entrada"

PREFIJO_TIPO = "Tipo de Documento:"
PREFIJO_GENERO = "Genero:"

# Artefactos del reporte bandeado. `Total` encabeza las 35 filas de subtotal por
# género y por tipo de documento; `Total General` es el cierre (con el total
# corrido a la 2da columna y los filtros del reporte en la 3ra y 4ta); `New Age
# Data` es el pie de página, y trae un número de página en la columna de
# unidades -- se ingiere como SKU fantasma con 1 unidad si no se filtra.
ARTEFACTOS = frozenset({"Total", "Total General", "New Age Data"})

# openpyxl devuelve los saltos de línea de una celda multilinea como este
# escape literal, no como \r\n.
SALTO_LITERAL = "_x000d__x000a_"


def parse_mes(label):
    """
    'Vta.Ago/25' o 'Ago/25' -> (2025, 8). None si no es una etiqueta de mes.

    La ventana se parsea del header en vez de hardcodearse: el archivo de hoy
    cubre Ago/25→Ago/26, pero el próximo que mande el cliente va a estar
    corrido y el módulo tiene que seguir andando sin tocarlo.
    """
    if not isinstance(label, str):
        return None
    txt = label.strip()
    if txt.startswith("Vta."):
        txt = txt[4:]
    if "/" not in txt:
        return None
    mes, _, anio = txt.partition("/")
    if mes not in MESES_ES or not anio.isdigit():
        return None
    return (2000 + int(anio) if len(anio) == 2 else int(anio), MESES_ES[mes])


def _vacio(valor):
    return valor is None or (isinstance(valor, str) and not valor.strip())


def normalizar_venta(valor):
    """
    Celda de venta -> int o None.

    None significa "no vendió", que es distinto de un cero: el cliente nunca
    escribe 0 en estas columnas. Un 0 explícito, si alguna vez apareciera, se
    respeta como 0 en vez de colapsarse a None -- son hechos distintos y el
    diagnóstico los tiene que poder distinguir.
    """
    if _vacio(valor):
        return None
    return int(valor)


def normalizar_rotacion(valor):
    """
    Celda de rotación -> float o None.

    Acá el cero SÍ es un cero real: las 13 columnas mensuales de rotación están
    densas (70.564 valores, todas las celdas) e incluyen ceros explícitos.
    """
    if _vacio(valor):
        return None
    return float(valor)


def normalizar_sku(valor):
    return str(valor).strip().upper()


def _texto(valor):
    """Limpia el espacio de no-quiebre que traen las descripciones del cliente."""
    if valor is None:
        return None
    return str(valor).replace("\xa0", " ").strip() or None


@dataclass
class ArticuloCliente:
    """
    Una fila del archivo de ventas, ya normalizada.

    Sin `frozen=True` a propósito: lleva dos dicts adentro, así que congelarlo
    prometería una inmutabilidad que no puede cumplir y además dejaría la clase
    sin `__hash__` utilizable (el sintetizado explota sobre los dicts).
    """
    sku: str
    descripcion: str | None
    barcode: str | None
    ventas: dict          # (year, month) -> int | None   (None = no vendió)
    rotaciones: dict      # (year, month) -> float | None (0.0 = cero real)
    rot_desestac: float | None
    rot_manual: float | None
    estado: str | None

    @property
    def tiene_venta(self):
        return any(v is not None for v in self.ventas.values())


def parsear_ventas(filas):
    """
    Filas del archivo de ventas -> {sku: ArticuloCliente}.

    Función pura: recibe la lista de tuplas que devuelve openpyxl, así que se
    testea sin leer ningún Excel.

    El layout se deduce del header, no de posiciones fijas. Las columnas con
    prefijo `Vta.` son venta y las que son sólo un mes son rotación mensual --
    es la convención del cliente, documentada en scripts/fixtures/
    planillas_cliente/README.md (#127).
    """
    filas = list(filas)
    if not filas:
        return {}

    header = filas[0]
    col_venta, col_rot, sueltas = {}, {}, {}
    for idx, etiqueta in enumerate(header):
        if not isinstance(etiqueta, str):
            continue
        limpia = etiqueta.strip()
        mes = parse_mes(limpia)
        if mes is not None:
            (col_venta if limpia.startswith("Vta.") else col_rot)[mes] = idx
        elif limpia in (COL_ROT_DESESTAC, COL_ROT_MANUAL, COL_ESTADO,
                        COL_DESCRIPCION, COL_BARCODE):
            sueltas[limpia] = idx

    def leer(fila, idx):
        return fila[idx] if idx is not None and idx < len(fila) else None

    articulos = {}
    for fila in filas[1:]:
        if not fila or _vacio(fila[0]):
            continue
        sku = normalizar_sku(fila[0])
        if sku in articulos:
            # Pisar en silencio perdería la venta de la primera fila y el
            # diagnóstico lo leería como "ese artículo vendió menos". En una
            # herramienta de diagnóstico, fallar fuerte es mejor que reportar
            # un número que nadie va a poder explicar después.
            raise ValueError(
                f"SKU duplicado en el archivo de ventas: {sku}. "
                "El cargador asume una fila por artículo; revisar el export.")
        articulos[sku] = ArticuloCliente(
            sku=sku,
            descripcion=_texto(leer(fila, sueltas.get(COL_DESCRIPCION))),
            barcode=_texto(leer(fila, sueltas.get(COL_BARCODE))),
            ventas={m: normalizar_venta(leer(fila, i)) for m, i in col_venta.items()},
            rotaciones={m: normalizar_rotacion(leer(fila, i)) for m, i in col_rot.items()},
            rot_desestac=normalizar_rotacion(leer(fila, sueltas.get(COL_ROT_DESESTAC))),
            rot_manual=normalizar_rotacion(leer(fila, sueltas.get(COL_ROT_MANUAL))),
            estado=_texto(leer(fila, sueltas.get(COL_ESTADO))),
        )
    return articulos


@dataclass(frozen=True)
class Movimiento:
    """Una fila de detalle del reporte de movimientos, con su banda resuelta."""
    sku: str
    descripcion: str | None
    unidades: int
    tipo_documento: str | None
    genero: str | None


@dataclass(frozen=True)
class FiltrosReporte:
    """
    Los filtros con los que el cliente generó el reporte, impresos en la fila
    `Total General`.

    `doc` es el que importa: una whitelist de tipos de documento. Acota lo que
    se puede concluir sobre #193 -- una transferencia entre depósitos con un
    código fuera de esta lista habría quedado excluida por el filtro, no por
    ausencia real, así que del archivo NO se puede deducir "no hay ninguna
    transferencia", sólo "no hay ninguno de los tipos pedidos".
    """
    lineas: tuple
    doc: tuple


def parsear_filtros(celdas):
    """Celdas de texto de la fila `Total General` -> FiltrosReporte."""
    lineas = []
    for celda in celdas:
        if not isinstance(celda, str):
            continue
        texto = celda.replace(SALTO_LITERAL, "\n").replace("\r\n", "\n")
        lineas.extend(l.strip() for l in texto.split("\n") if l.strip())

    doc = ()
    for linea in lineas:
        if linea.startswith("Doc.:"):
            doc = tuple(d.strip() for d in linea[len("Doc.:"):].split(",") if d.strip())
            break
    return FiltrosReporte(lineas=tuple(lineas), doc=doc)


def _como_entero(valor):
    """
    Unidades -> int, o None si la celda no es un número.

    Pasa por `float` antes de `int` porque una celda puede llegar como '12.0':
    validar con `float` y convertir con `int` directamente hace que esa celda
    pase el filtro y después reviente la corrida entera.
    """
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return None


def parsear_movimientos(filas):
    """
    Filas del reporte de movimientos -> ([Movimiento], FiltrosReporte).

    Función pura. El reporte es bandeado: las filas `Tipo de Documento: X` y
    `Genero: Y` son secciones que hay que arrastrar como estado mientras se
    recorre, no datos.
    """
    movimientos, celdas_filtro = [], []
    tipo = genero = None

    # Se recorren todas las filas, incluida la primera: la de header cae sola
    # porque su columna de unidades dice ' Unidades' y no es un número. Saltear
    # la primera a ciegas rompería un export que arranque con una banda en vez
    # de con el header -- ahí todas las filas de esa banda quedarían sin tipo de
    # documento y desaparecerían del emparejamiento sin aviso.
    for fila in list(filas):
        if not fila or _vacio(fila[0]):
            continue
        primera = str(fila[0]).strip()

        if primera.startswith(PREFIJO_TIPO):
            tipo = primera[len(PREFIJO_TIPO):].strip()
            continue
        if primera.startswith(PREFIJO_GENERO):
            genero = primera[len(PREFIJO_GENERO):].strip()
            continue
        if primera in ARTEFACTOS:
            # `Total General` trae los filtros del reporte en las columnas 3 y
            # 4 -- es la única fila artefacto de la que se rescata algo.
            if primera == "Total General":
                celdas_filtro.extend(fila[2:])
            continue

        unidades = _como_entero(fila[2] if len(fila) > 2 else None)
        if unidades is None:
            continue

        movimientos.append(Movimiento(
            sku=normalizar_sku(primera),
            descripcion=_texto(fila[1] if len(fila) > 1 else None),
            unidades=unidades,
            tipo_documento=tipo,
            genero=genero,
        ))

    return movimientos, parsear_filtros(celdas_filtro)


@dataclass(frozen=True)
class Recodificacion:
    """Una Entrada de Mercadería y las salidas del mismo mes que la explican."""
    destino: str
    origenes: tuple
    unidades: int


def _subset_que_suma(candidatos, objetivo):
    """
    Devuelve los índices de un subconjunto de `candidatos` (enteros positivos)
    que suma exactamente `objetivo`, o None.

    DP de subset-sum. Los montos del reporte son enteros chicos y las salidas
    son decenas de filas, así que alcanza de sobra. Se prefiere el subconjunto
    más chico para no arrastrar salidas de más a un emparejamiento que ya cierra
    con menos.
    """
    if objetivo <= 0:
        return None
    # alcanzables[suma] = índices que la producen, con la menor cantidad posible
    alcanzables = {0: ()}
    for idx, valor in enumerate(candidatos):
        if valor <= 0 or valor > objetivo:
            continue
        for suma, indices in sorted(alcanzables.items(), reverse=True):
            nueva = suma + valor
            if nueva > objetivo:
                continue
            if nueva not in alcanzables or len(indices) + 1 < len(alcanzables[nueva]):
                alcanzables[nueva] = indices + (idx,)
    return alcanzables.get(objetivo)


_RUIDO = frozenset({"NO", "USAR", "DE", "LA", "EL", "CON", "SIN", "PARA"})

# Jaccard mínimo entre descripciones para considerar que dos filas son el mismo
# producto. Bajo a propósito: las descripciones de un par real difieren en el
# prefijo "(NO USAR)", en el código de modelo o en el orden de las palabras
# ("Carcasa A CHUWI Corebook X I5 repuesto" vs "Carcasa A Bisagra Corebook X
# repuesto CHUWI"), pero comparten el grueso. Sólo tiene que separar un par real
# de uno que coincide por monto y nada más.
UMBRAL_PARECIDO = 0.2


def _tokens(texto):
    if not texto:
        return frozenset()
    limpio = "".join(c if c.isalnum() else " " for c in texto.upper())
    return frozenset(t for t in limpio.split() if len(t) >= 3 and t not in _RUIDO)


def _similitud(a, b):
    ta, tb = _tokens(a), _tokens(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


def emparejar_recodificaciones(movimientos):
    """
    Empareja cada Entrada de Mercadería contra las salidas del mismo mes que la
    explican, por monto y sin lista hardcodeada.

    En el archivo real el 100% de las +466 unidades de Entrada de Mercadería es
    la contrapartida exacta de 466 de las 819 unidades de Ajuste de Salida: es
    renumeración de códigos, no stock perdido. Sin este emparejamiento, códigos
    de promo retirados como HX030/HX031 se reportan como huecos de nuestro
    catálogo cuando en realidad sus unidades volvieron bajo SKUs que sí tenemos.

    El monto solo no alcanza para desempatar: en el archivo real hay dos salidas
    de 14 unidades y sólo una es la contrapartida de K00012 (+14). Por eso se
    resuelve en dos pasadas -- primero contra las salidas que además comparten
    identidad de producto por descripción, y recién después contra el resto. Sin
    eso el par sale armado con un cargador FONENG en lugar de una carcasa CHUWI,
    y el chequeo de continuidad de stock de #199 mira la serie equivocada.

    Una salida no se reutiliza en dos entradas. Las entradas se resuelven de
    menor a mayor para que las que cierran con una sola salida se lleven su par
    antes de que una entrada grande se lo coma dentro de un subconjunto.
    """
    salidas = [m for m in movimientos if m.tipo_documento == AJUSTE_SALIDA and m.unidades < 0]
    entradas = sorted((m for m in movimientos if m.tipo_documento == ENTRADA_MERCADERIA),
                      key=lambda m: m.unidades)

    disponibles = list(range(len(salidas)))
    pares = []
    for entrada in entradas:
        parecidas = [i for i in disponibles
                     if _similitud(entrada.descripcion, salidas[i].descripcion) >= UMBRAL_PARECIDO]

        indices = None
        for candidatas in (parecidas, disponibles):
            if not candidatas:
                continue
            elegidos = _subset_que_suma([abs(salidas[i].unidades) for i in candidatas],
                                        entrada.unidades)
            if elegidos is not None:
                indices = [candidatas[i] for i in elegidos]
                break
        if indices is None:
            continue

        pares.append(Recodificacion(
            destino=entrada.sku,
            origenes=tuple(salidas[i].sku for i in indices),
            unidades=entrada.unidades,
        ))
        disponibles = [i for i in disponibles if i not in set(indices)]
    return tuple(pares)


@dataclass(frozen=True)
class Cobertura:
    comunes: tuple
    solo_cliente: tuple
    solo_nuestros: tuple


def comparar_catalogo(skus_cliente, skus_nuestros):
    """Cruza ambos catálogos por SKU normalizado."""
    del_cliente = {normalizar_sku(s) for s in skus_cliente}
    nuestros = {normalizar_sku(s) for s in skus_nuestros}
    return Cobertura(
        comunes=tuple(sorted(del_cliente & nuestros)),
        solo_cliente=tuple(sorted(del_cliente - nuestros)),
        solo_nuestros=tuple(sorted(nuestros - del_cliente)),
    )


@dataclass(frozen=True)
class AjusteCobertura:
    """
    Parte los SKUs ausentes de `articulos` en los que son un hueco real y los
    que son un código que el cliente retiró.
    """
    ausencias_reales: tuple
    recodificados: tuple
    pares: tuple


def ajustar_cobertura_por_recodificacion(cobertura, movimientos):
    """
    Un SKU ausente de nuestro catálogo que aparece como origen de una
    recodificación no es un hueco nuestro: es un código que el cliente retiró y
    cuyas unidades volvieron bajo otro SKU.

    En el archivo real esto saca a HX030 y HX031 de la lista de huecos -- entre
    los dos son 350 de las 819 unidades de Ajuste de Salida (43%), así que
    reportarlos como ausencia nuestra sería el error más caro del ticket.
    """
    pares = emparejar_recodificaciones(movimientos)
    origenes = {sku for p in pares for sku in p.origenes}
    ausentes = set(cobertura.solo_cliente)
    return AjusteCobertura(
        ausencias_reales=tuple(sorted(ausentes - origenes)),
        recodificados=tuple(sorted(ausentes & origenes)),
        pares=pares,
    )


# ── lectura de los archivos ──────────────────────────────────────────────────

def _filas(path, hoja):
    import openpyxl  # diferido a propósito, ver nota del encabezado

    libro = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if hoja in libro.sheetnames:
        nombre = hoja
    elif len(libro.sheetnames) == 1:
        # Un export de una sola hoja con otro nombre es benigno (los dos
        # archivos del cliente difieren entre sí: 'Ventas' y 'Sheet1').
        nombre = libro.sheetnames[0]
    else:
        # Con varias hojas, elegir la primera a ciegas produce una salida
        # plausible pero sin sentido, y nadie se entera.
        raise ValueError(
            f"{os.path.basename(path)} no tiene la hoja '{hoja}'. "
            f"Hojas disponibles: {', '.join(libro.sheetnames)}")
    return list(libro[nombre].iter_rows(values_only=True))


def leer_ventas(path=None, hoja=HOJA_VENTAS):
    return parsear_ventas(_filas(path or os.path.join(DIR_POR_DEFECTO, ARCHIVO_VENTAS), hoja))


def leer_movimientos(path=None, hoja=HOJA_MOVIMIENTOS):
    return parsear_movimientos(
        _filas(path or os.path.join(DIR_POR_DEFECTO, ARCHIVO_MOVIMIENTOS), hoja))


def leer_skus_nuestros(conn=None):
    """
    SKUs de `articulos`. Devuelve None si no hay credenciales configuradas y no
    se pasó una conexión ya abierta.

    Acepta una conexión existente (`conn`) para que un script que ya mantiene
    la suya para varias queries -- como comparar_ventas_mensual.py -- no tenga
    que abrir una segunda en paralelo sólo para este SELECT.
    """
    if conn is not None:
        with conn.cursor() as cur:
            cur.execute("SELECT sku FROM articulos")
            return [f[0] for f in cur.fetchall()]

    if not os.environ.get("MYSQL_PASSWORD"):
        return None
    import pymysql

    conn = pymysql.connect(
        host=os.environ.get("MYSQL_HOST", "mysql"),
        port=int(os.environ.get("MYSQL_PORT", "3306")),
        user=os.environ["MYSQL_USER"],
        password=os.environ["MYSQL_PASSWORD"],
        database=os.environ.get("MYSQL_DB", "evalutia"),
    )
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT sku FROM articulos")
            return [f[0] for f in cur.fetchall()]
    finally:
        conn.close()


# ── informe ──────────────────────────────────────────────────────────────────

def main():
    ventas = leer_ventas()
    movimientos, filtros = leer_movimientos()

    print("=" * 78)
    print("ARCHIVO 1 -- ventas")
    print("=" * 78)
    meses = sorted({m for a in ventas.values() for m in a.ventas})
    print(f"  SKUs: {len(ventas)}")
    if meses:
        print(f"  Ventana: {len(meses)} meses, {meses[0]} -> {meses[-1]}")
    else:
        print("  Ventana: (ninguna columna de mes reconocida en el header)")
    celdas = [v for a in ventas.values() for v in a.ventas.values()]
    print(f"  Celdas de venta: {len(celdas)} "
          f"(nulas {sum(1 for v in celdas if v is None)}, "
          f"cero explícito {sum(1 for v in celdas if v == 0)}, "
          f"negativas {sum(1 for v in celdas if v is not None and v < 0)}, "
          f"positivas {sum(1 for v in celdas if v is not None and v > 0)})")
    print(f"  SKUs con al menos un mes de venta: "
          f"{sum(1 for a in ventas.values() if a.tiene_venta)}")
    print(f"  Con código de barras: {sum(1 for a in ventas.values() if a.barcode)}")
    con_rot = [a for a in ventas.values() if a.rot_desestac is not None]
    iguales = sum(1 for a in con_rot if a.rot_manual == a.rot_desestac)
    print(f"  Con rotación cargada: {len(con_rot)} "
          f"(Rot. Manual == Rotacion DesEstac. en {iguales})")
    if con_rot and iguales == len(con_rot):
        print("    -> es el default del sistema, no un override manual: en los archivos")
        print("       `bruto` de #127 la columna viene precargada con Rotacion DesEstac.")
        print("       y la persona que arma el pedido decide apartarse de ella. Este")
        print("       export no trae ninguna decisión experta cargada.")
    estados = {}
    for a in ventas.values():
        estados[a.estado] = estados.get(a.estado, 0) + 1
    print(f"  Estado Art.: {estados}")

    print()
    print("=" * 78)
    print("ARCHIVO 2 -- movimientos de stock")
    print("=" * 78)
    por_tipo = {}
    for m in movimientos:
        por_tipo[m.tipo_documento] = por_tipo.get(m.tipo_documento, 0) + m.unidades
    print(f"  Filas de detalle: {len(movimientos)}  |  SKUs distintos: "
          f"{len({m.sku for m in movimientos})}")
    for tipo, total in sorted(por_tipo.items(), key=lambda kv: str(kv[0])):
        print(f"    {tipo or '(sin tipo de documento)'}: {total:+d}")
    print(f"  Neto: {sum(m.unidades for m in movimientos):+d}")
    print(f"  Géneros: {len({m.genero for m in movimientos})}")
    print()
    print("  Filtros con los que se generó el reporte:")
    for linea in filtros.lineas:
        print(f"    {linea}")
    print(f"  -> whitelist de documentos: {', '.join(filtros.doc) or '(ninguna)'}")
    print("     OJO: una transferencia entre depósitos con un código fuera de esa")
    print("     lista quedó excluida POR EL FILTRO, no por ausencia real. De este")
    print("     archivo no se puede concluir 'no hay transferencias' (#193/#199).")

    pares = emparejar_recodificaciones(movimientos)
    total_emr = sum(m.unidades for m in movimientos if m.tipo_documento == ENTRADA_MERCADERIA)
    print()
    print(f"  Recodificaciones de SKU: {len(pares)} pares, "
          f"{sum(p.unidades for p in pares)} de {total_emr} unidades de entrada")
    for p in pares:
        print(f"    {' + '.join(p.origenes):<28} -> {p.destino:<10} {p.unidades:+d}")

    print()
    print("=" * 78)
    print("COBERTURA DE CATÁLOGO")
    print("=" * 78)
    nuestros = leer_skus_nuestros()
    if nuestros is None:
        print("  (sin credenciales de MySQL: se omite el cruce contra `articulos`)")
        return 0

    print(f"  Artículos en `articulos`: {len(nuestros)}")
    for nombre, skus in (("ventas", set(ventas)),
                         ("movimientos", {m.sku for m in movimientos})):
        cob = comparar_catalogo(skus, nuestros)
        print(f"\n  Archivo de {nombre}: {len(skus)} SKUs")
        print(f"    comunes con `articulos`: {len(cob.comunes)}")
        print(f"    suyos que no tenemos:    {len(cob.solo_cliente)}")
        if nombre == "movimientos":
            ajuste = ajustar_cobertura_por_recodificacion(cob, movimientos)
            print(f"      de los cuales recodificados (NO son hueco nuestro): "
                  f"{', '.join(ajuste.recodificados) or '(ninguno)'}")
            print(f"      ausencias reales: "
                  f"{', '.join(ajuste.ausencias_reales) or '(ninguna)'}")
        elif cob.solo_cliente:
            muestra = ', '.join(cob.solo_cliente[:15])
            print(f"      {muestra}{' ...' if len(cob.solo_cliente) > 15 else ''}")
        print(f"    nuestros que no están en su archivo: {len(cob.solo_nuestros)}")

    comun = comparar_catalogo(set(ventas), nuestros).comunes
    print(f"\n  Conjunto común para #196 en adelante: {len(comun)} SKUs")
    return 0


if __name__ == "__main__":
    sys.exit(main())
